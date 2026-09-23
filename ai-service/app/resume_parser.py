"""Extracts text and structured signals (skills, experience, sections, quality) from resumes."""
import io
import re
from datetime import date

from docx import Document
from pypdf import PdfReader

from .skills import extract_skills

SECTION_KEYWORDS = {
    "Summary": ["summary", "objective", "profile", "about me"],
    "Experience": ["experience", "work history", "employment", "internship", "internships"],
    "Education": ["education", "academic", "qualification"],
    "Skills": ["skills", "technical skills", "technologies", "tech stack"],
    "Projects": ["projects", "personal projects", "academic projects"],
    "Certifications": ["certifications", "certificates", "courses"],
    "Achievements": ["achievements", "awards", "accomplishments", "honors"],
}

ACTION_VERBS = {"built", "developed", "designed", "implemented", "led", "created", "optimized", "improved",
                "reduced", "increased", "deployed", "automated", "architected", "delivered", "launched",
                "migrated", "integrated", "managed", "mentored", "engineered"}

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(?:\+?\d{1,3}[\s-]?)?(?:\(?\d{3,5}\)?[\s-]?)\d{3,5}[\s-]?\d{3,5}")
YEARS_RE = re.compile(r"(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years?|yrs?)(?:\s+of)?\s+(?:\w+\s+){0,3}?experience",
                      re.IGNORECASE)
RANGE_RE = re.compile(r"((?:19|20)\d{2})\s*(?:-|–|—|to)\s*((?:19|20)\d{2}|present|current|now)", re.IGNORECASE)


class UnsupportedFileError(ValueError):
    pass


def extract_text(content: bytes, filename: str) -> str:
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf"):
            reader = PdfReader(io.BytesIO(content))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif name.endswith(".docx"):
            doc = Document(io.BytesIO(content))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.extend(cell.text for cell in row.cells)
            text = "\n".join(parts)
        elif name.endswith(".txt"):
            text = content.decode("utf-8", errors="ignore")
        else:
            raise UnsupportedFileError("Only PDF, DOCX or TXT resumes are supported")
    except UnsupportedFileError:
        raise
    except Exception as exc:  # corrupt / encrypted files
        raise UnsupportedFileError(f"Could not read the file: {exc}") from exc
    # collapse runs of spaces but keep line breaks (section detection relies on them)
    return "\n".join(re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines() if line.strip())


def detect_sections(text: str) -> list[str]:
    found = []
    lines = [line.strip().lower().rstrip(":") for line in text.splitlines()]
    for section, keywords in SECTION_KEYWORDS.items():
        # a heading is a short line that starts with the keyword
        if any(len(line) <= 40 and any(line.startswith(k) for k in keywords) for line in lines):
            found.append(section)
    return found


def estimate_experience_years(text: str) -> int | None:
    explicit = [float(m) for m in YEARS_RE.findall(text) if float(m) < 45]
    if explicit:
        return int(max(explicit))
    total = 0
    this_year = date.today().year
    for start, end in RANGE_RE.findall(text):
        end_year = this_year if end.lower() in {"present", "current", "now"} else int(end)
        if 0 <= end_year - int(start) <= 45:
            total += end_year - int(start)
    # year ranges also cover education, so only trust them when an experience section exists
    if total and "Experience" in detect_sections(text):
        return min(total, 40)
    return None


def assess_quality(text: str, sections: list[str], skills: list[str]) -> tuple[int, list[str]]:
    """Heuristic ATS-readiness score (0-100) with actionable tips."""
    score = 0
    tips: list[str] = []
    words = re.findall(r"[A-Za-z]+", text)
    lower_words = {w.lower() for w in words}

    score += min(len(sections), 6) * 7  # up to 42
    for must in ("Experience", "Education", "Skills", "Projects"):
        if must not in sections:
            tips.append(f"Add a clearly titled '{must}' section so ATS parsers can find it.")

    if EMAIL_RE.search(text):
        score += 8
    else:
        tips.append("Include a professional email address.")
    if PHONE_RE.search(text):
        score += 5
    else:
        tips.append("Include a phone number.")

    if 250 <= len(words) <= 1200:
        score += 15
    elif len(words) < 250:
        score += 5
        tips.append("Your resume is short. Describe projects and impact in more detail (aim for 400-800 words).")
    else:
        score += 8
        tips.append("Your resume is long. Keep it to 1-2 pages focused on relevant experience.")

    quantified = len(re.findall(r"\d+(?:\.\d+)?\s*(?:%|x\b|\+|k\b|users|ms\b|requests)", text, re.IGNORECASE))
    score += min(quantified, 5) * 3  # up to 15
    if quantified < 2:
        tips.append("Quantify achievements (e.g. 'reduced API latency by 40%', 'served 10k users').")

    verbs = len(ACTION_VERBS & lower_words)
    score += min(verbs, 5) * 1  # up to 5
    if verbs < 3:
        tips.append("Start bullet points with strong action verbs like 'Built', 'Designed', 'Optimized'.")

    score += min(len(skills), 10)  # up to 10
    if len(skills) < 6:
        tips.append("List more relevant technical skills (languages, frameworks, tools).")

    return min(score, 100), tips


def parse_resume(content: bytes, filename: str) -> dict:
    text = extract_text(content, filename)
    skills = extract_skills(text)
    sections = detect_sections(text)
    quality, tips = assess_quality(text, sections, skills)
    email = EMAIL_RE.search(text)
    phone = PHONE_RE.search(text)
    return {
        "text": text,
        "skills": skills,
        "experience_years": estimate_experience_years(text),
        "email": email.group(0) if email else None,
        "phone": phone.group(0).strip() if phone else None,
        "sections_found": sections,
        "quality_score": quality,
        "tips": tips,
    }
